# 開発用: 書き出した xlsx を元ファイルと比較する
# 使い方: python verify.py <元.xlsx> <出力.xlsx> <列名> [シート名] [前回の列名]
import sys
import zipfile
from openpyxl import load_workbook

src, out, col = sys.argv[1], sys.argv[2], sys.argv[3]
sheet = sys.argv[4] if len(sys.argv) > 4 else None
prev_col = sys.argv[5] if len(sys.argv) > 5 else "G"

# 1. シート以外の ZIP エントリはバイト単位で同一か
za, zb = zipfile.ZipFile(src), zipfile.ZipFile(out)
for n in za.namelist():
    if "worksheets/" in n:
        continue
    assert za.read(n) == zb.read(n), f"changed: {n}"
assert sorted(za.namelist()) == sorted(zb.namelist()), "entry list differs"
assert zb.testzip() is None
print("zip entries other than sheets: identical")

wa, wb = load_workbook(src), load_workbook(out)
a, b = (wa[sheet], wb[sheet]) if sheet else (wa.active, wb.active)

# 2. 対象列以外は変化なし
for row in a.iter_rows(min_row=1, max_row=a.max_row, max_col=a.max_column):
    for c in row:
        if getattr(c, "column_letter", None) == col:
            continue  # 書き込み対象の列は 3. で確認する
        d = b[c.coordinate]
        assert c.value == d.value and c.number_format == d.number_format, f"diff {c.coordinate}"
assert [str(m) for m in a.merged_cells.ranges] == [str(m) for m in b.merged_cells.ranges]
print("other cells/merges: unchanged")

# 3. 対象列の中身
for r in range(1, b.max_row + 1):
    c = b[f"{col}{r}"]
    prev = b[f"{prev_col}{r}"]
    if c.value is None and prev.value is None:
        continue
    print(f"{col}{r:<3} {c.value!r:<28} fmt={c.number_format:<10} | {prev_col}: {prev.value!r} fmt={prev.number_format}")
print("dimension:", b.dimensions)
